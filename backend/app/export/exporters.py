"""PDF / DOCX export (§50). Preserves layout, typography, section hierarchy."""
from __future__ import annotations

import io
from pathlib import Path

from ..evidence_citations import strip_inline_evidence

_FONTS_DIR = Path(__file__).resolve().parent / "fonts"

_HEADINGS = {
    "en": {
        "summary": "PROFESSIONAL SUMMARY",
        "experience": "PROFESSIONAL EXPERIENCE",
        "skills": "SKILLS",
        "education": "EDUCATION",
        "certifications": "CERTIFICATIONS",
        "languages": "LANGUAGES",
        "additional": "ADDITIONAL",
        "salutation": "Dear Hiring Manager,",
        "closing": "Sincerely,",
    },
    "ar": {
        "summary": "الملخص المهني",
        "experience": "الخبرة المهنية",
        "skills": "المهارات",
        "education": "التعليم",
        "certifications": "الشهادات",
        "languages": "اللغات",
        "additional": "إضافي",
        "salutation": "السيد/ة مدير التوظيف،",
        "closing": "مع فائق الاحترام،",
    },
}


def _headings(language: str) -> dict[str, str]:
    return _HEADINGS["ar"] if language == "ar" else _HEADINGS["en"]


def _shape(text: str, language: str) -> str:
    if language != "ar" or not text:
        return text
    try:
        import arabic_reshaper
        try:
            from bidi.algorithm import get_display
        except ImportError:
            from bidi import get_display
        return get_display(arabic_reshaper.reshape(text))
    except Exception:
        return text


def _arabic_font_paths() -> tuple[Path | None, Path | None]:
    regular = _FONTS_DIR / "IBMPlexSansArabic-Regular.ttf"
    bold = _FONTS_DIR / "IBMPlexSansArabic-Bold.ttf"
    return (regular if regular.is_file() else None, bold if bold.is_file() else None)


def _register_arabic_fonts() -> tuple[str, str] | None:
    regular, bold = _arabic_font_paths()
    if regular is None:
        return None
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    if "IBMPlexSansArabic" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("IBMPlexSansArabic", str(regular)))
        if bold is not None:
            pdfmetrics.registerFont(TTFont("IBMPlexSansArabic-Bold", str(bold)))
    return ("IBMPlexSansArabic", "IBMPlexSansArabic-Bold" if bold else "IBMPlexSansArabic")


def _visible_text(text: str | None) -> str:
    return strip_inline_evidence(text or "")[0]


def _hex_to_rgb(hex_color: str) -> tuple[float, float, float]:
    hex_color = hex_color.lstrip("#")
    if len(hex_color) != 6:
        return (0.0, 0.0, 0.0)
    return tuple(int(hex_color[i : i + 2], 16) / 255 for i in (0, 2, 4))


def export_pdf(document_type: str, content: dict, profile: dict, template_config: dict, content_language: str = "en") -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer

    typo = template_config.get("typography", {})
    spacing = template_config.get("spacing", {})
    colors = template_config.get("colors", {})
    layout = template_config.get("layout", {})

    buffer = io.BytesIO()
    margin = spacing.get("margins", 54)
    doc = SimpleDocTemplate(
        buffer, pagesize=A4,
        leftMargin=margin, rightMargin=margin, topMargin=margin, bottomMargin=margin,
        title=profile.get("full_name", "Document"),
    )

    body_font = typo.get("body_font", "Helvetica")
    heading_font = typo.get("heading_font", "Helvetica-Bold")
    arabic_fonts = _register_arabic_fonts() if content_language == "ar" else None
    if arabic_fonts:
        body_font, heading_font = arabic_fonts
    body_size = typo.get("body_size", 10)
    heading_size = typo.get("heading_size", 13)
    name_size = typo.get("name_size", 20)
    accent = _hex_to_rgb(colors.get("accent", "#2d6cdf"))
    heading_color = _hex_to_rgb(colors.get("heading", "#1a1a2e"))
    body_color = _hex_to_rgb(colors.get("body", "#222222"))
    muted = _hex_to_rgb(colors.get("muted", "#666666"))
    section_gap = spacing.get("section_gap", 14)
    entry_gap = spacing.get("entry_gap", 8)

    from reportlab.lib.enums import TA_LEFT, TA_RIGHT

    align = TA_RIGHT if content_language == "ar" else TA_LEFT
    lang = "ar" if content_language == "ar" else "en"
    heads = _headings(lang)

    def tx(value: str) -> str:
        return _shape(value, lang)

    name_style = ParagraphStyle(
        "name",
        fontName=heading_font,
        fontSize=name_size,
        leading=name_size * 1.25,
        textColor=heading_color,
        spaceAfter=8,
        alignment=align,
    )
    headline_style = ParagraphStyle(
        "headline",
        fontName=body_font,
        fontSize=body_size + 1,
        leading=(body_size + 1) * 1.3,
        textColor=muted,
        spaceAfter=2,
        alignment=align,
    )
    h_style = ParagraphStyle(
        "h",
        fontName=heading_font,
        fontSize=heading_size,
        leading=heading_size * 1.25,
        textColor=accent,
        spaceBefore=section_gap,
        spaceAfter=4,
        alignment=align,
    )
    b_style = ParagraphStyle("b", fontName=body_font, fontSize=body_size, textColor=body_color, leading=body_size * 1.35, alignment=align)
    bullet_style = ParagraphStyle("bullet", parent=b_style, leftIndent=spacing.get("bullet_indent", 12), bulletIndent=0, spaceAfter=2)
    entry_style = ParagraphStyle("entry", fontName=heading_font, fontSize=body_size + 1, textColor=heading_color, spaceBefore=entry_gap, alignment=align)
    date_style = ParagraphStyle("date", fontName=body_font, fontSize=body_size - 1, textColor=muted, spaceAfter=2, alignment=align)

    story: list = []
    bullet_char = layout.get("bullet_char", "•")

    def rule():
        if layout.get("header_rule"):
            story.append(HRFlowable(width="100%", thickness=0.7, color=accent, spaceBefore=2, spaceAfter=6))

    if document_type == "cv":
        if profile.get("full_name"):
            story.append(Paragraph(tx(profile["full_name"]), name_style))
        if content.get("headline"):
            story.append(Paragraph(tx(content["headline"]), headline_style))
        contact = " | ".join(p for p in [profile.get("email", ""), profile.get("phone", ""), profile.get("location", "")] if p)
        if contact:
            story.append(Paragraph(tx(contact), date_style))
        rule()

        if content.get("summary"):
            story.append(Paragraph(tx(heads["summary"]), h_style))
            story.append(Paragraph(tx(_visible_text(content["summary"])), b_style))

        if content.get("experience"):
            story.append(Paragraph(tx(heads["experience"]), h_style))
            for entry in content["experience"]:
                header = f"{entry.get('job_title', '')}, {entry.get('organization', '')}"
                story.append(Paragraph(tx(header), entry_style))
                dates = " – ".join(p for p in [entry.get("start_date", ""), entry.get("end_date", "")] if p)
                loc = entry.get("location", "")
                meta = " | ".join(p for p in [dates, loc] if p)
                if meta:
                    story.append(Paragraph(tx(meta), date_style))
                for bullet in entry.get("bullets", []):
                    text = bullet.get("text", bullet) if isinstance(bullet, dict) else str(bullet)
                    visible = _visible_text(str(text))
                    if visible:
                        story.append(Paragraph(tx(visible), bullet_style, bulletText=bullet_char))

        if content.get("skills"):
            story.append(Paragraph(tx(heads["skills"]), h_style))
            story.append(Paragraph(tx(", ".join(content["skills"])), b_style))

        if content.get("education"):
            story.append(Paragraph(tx(heads["education"]), h_style))
            for ed in content["education"]:
                parts = [ed.get("degree", ""), ed.get("field_of_study", ""), ed.get("institution", "")]
                story.append(Paragraph(tx(", ".join(p for p in parts if p)), entry_style))
                dates = " – ".join(p for p in [ed.get("start_date", ""), ed.get("end_date", "")] if p)
                if dates:
                    story.append(Paragraph(tx(dates), date_style))

        if content.get("certifications"):
            story.append(Paragraph(tx(heads["certifications"]), h_style))
            for cert in content["certifications"]:
                story.append(Paragraph(tx(str(cert)), bullet_style, bulletText=bullet_char))

        if content.get("languages"):
            story.append(Paragraph(tx(heads["languages"]), h_style))
            story.append(Paragraph(
                tx(", ".join(f"{l.get('language', '')} ({l.get('proficiency', '')})".strip(" ()") for l in content["languages"])),
                b_style,
            ))

        for section in content.get("additional_sections", []):
            story.append(Paragraph(tx(section.get("name", heads["additional"]).upper()), h_style))
            for entry in section.get("entries", []):
                text = entry.get("text", entry) if isinstance(entry, dict) else str(entry)
                story.append(Paragraph(tx(str(text)), bullet_style, bulletText=bullet_char))
    else:
        if profile.get("full_name"):
            story.append(Paragraph(tx(profile["full_name"]), name_style))
        contact = " | ".join(p for p in [profile.get("email", ""), profile.get("phone", ""), profile.get("location", "")] if p)
        if contact:
            story.append(Paragraph(tx(contact), date_style))
        story.append(Spacer(1, 8 * mm))
        if content.get("subject"):
            story.append(Paragraph(f"<b>{tx(content['subject'])}</b>", b_style))
            story.append(Spacer(1, 4 * mm))
        story.append(Paragraph(tx(content.get("salutation", heads["salutation"])), b_style))
        story.append(Spacer(1, 3 * mm))
        for para in content.get("paragraphs", []):
            text = para.get("text", "") if isinstance(para, dict) else str(para)
            visible = _visible_text(str(text))
            if visible:
                story.append(Paragraph(tx(visible), b_style))
                story.append(Spacer(1, 3 * mm))
        story.append(Spacer(1, 4 * mm))
        for line in content.get("closing", heads["closing"]).split("\n"):
            story.append(Paragraph(tx(line), b_style))

    doc.build(story)
    return buffer.getvalue()


def export_docx(document_type: str, content: dict, profile: dict, template_config: dict, content_language: str = "en") -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor

    typo = template_config.get("typography", {})
    colors = template_config.get("colors", {})
    layout = template_config.get("layout", {})

    document = Document()
    style = document.styles["Normal"]
    style.font.name = typo.get("body_font", "Helvetica").replace("-Roman", "").replace("-Bold", "")
    style.font.size = Pt(typo.get("body_size", 10))

    accent = _hex_to_rgb(colors.get("accent", "#2d6cdf"))
    accent_rgb = RGBColor(*(int(c * 255) for c in accent))
    bullet_char = layout.get("bullet_char", "•")

    lang = "ar" if content_language == "ar" else "en"
    heads = _headings(lang)

    def heading(text: str):
        p = document.add_paragraph()
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(typo.get("heading_size", 13))
        run.font.color.rgb = accent_rgb
        return p

    def bullet(text: str):
        return document.add_paragraph(f"{bullet_char} {text}" if bullet_char else text)

    if document_type == "cv":
        name = document.add_paragraph()
        name.paragraph_format.space_after = Pt(8)
        run = name.add_run(profile.get("full_name", ""))
        run.bold = True
        run.font.size = Pt(typo.get("name_size", 20))
        if content.get("headline"):
            document.add_paragraph(content["headline"])
        contact = " | ".join(p for p in [profile.get("email", ""), profile.get("phone", ""), profile.get("location", "")] if p)
        if contact:
            document.add_paragraph(contact)

        if content.get("summary"):
            heading(heads["summary"])
            document.add_paragraph(_visible_text(content["summary"]))
        if content.get("experience"):
            heading(heads["experience"])
            for entry in content["experience"]:
                p = document.add_paragraph()
                r = p.add_run(f"{entry.get('job_title', '')}, {entry.get('organization', '')}")
                r.bold = True
                dates = " – ".join(x for x in [entry.get("start_date", ""), entry.get("end_date", "")] if x)
                meta = " | ".join(x for x in [dates, entry.get("location", "")] if x)
                if meta:
                    document.add_paragraph(meta)
                for b in entry.get("bullets", []):
                    visible = _visible_text(str(b.get("text", b) if isinstance(b, dict) else b))
                    if visible:
                        bullet(visible)
        if content.get("skills"):
            heading(heads["skills"])
            document.add_paragraph(", ".join(content["skills"]))
        if content.get("education"):
            heading(heads["education"])
            for ed in content["education"]:
                parts = [ed.get("degree", ""), ed.get("field_of_study", ""), ed.get("institution", "")]
                document.add_paragraph(", ".join(p for p in parts if p))
        if content.get("certifications"):
            heading(heads["certifications"])
            for cert in content["certifications"]:
                bullet(cert)
        if content.get("languages"):
            heading(heads["languages"])
            document.add_paragraph(
                ", ".join(f"{l.get('language', '')} ({l.get('proficiency', '')})".strip(" ()") for l in content["languages"])
            )
        for section in content.get("additional_sections", []):
            heading(section.get("name", heads["additional"]).upper())
            for entry in section.get("entries", []):
                bullet(entry.get("text", entry) if isinstance(entry, dict) else str(entry))
    else:
        if profile.get("full_name"):
            p = document.add_paragraph()
            r = p.add_run(profile["full_name"])
            r.bold = True
            r.font.size = Pt(typo.get("name_size", 16))
        contact = " | ".join(p for p in [profile.get("email", ""), profile.get("phone", ""), profile.get("location", "")] if p)
        if contact:
            document.add_paragraph(contact)
        document.add_paragraph("")
        if content.get("subject"):
            p = document.add_paragraph()
            p.add_run(content["subject"]).bold = True
        document.add_paragraph(content.get("salutation", heads["salutation"]))
        for para in content.get("paragraphs", []):
            text = para.get("text", "") if isinstance(para, dict) else str(para)
            visible = _visible_text(str(text))
            if visible:
                document.add_paragraph(visible)
        for line in content.get("closing", heads["closing"]).split("\n"):
            document.add_paragraph(line)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def export_document(
    fmt: str,
    document_type: str,
    content: dict,
    profile: dict,
    template_config: dict,
    content_language: str = "en",
) -> tuple[bytes, str]:
    if fmt == "pdf":
        return export_pdf(document_type, content, profile, template_config, content_language), "application/pdf"
    if fmt == "docx":
        return export_docx(document_type, content, profile, template_config, content_language), (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
    raise ValueError(f"Unsupported export format: {fmt}")
